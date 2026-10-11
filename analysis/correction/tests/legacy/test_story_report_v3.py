from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Published paper\high-order-topology\Data_processing\Result")
REPORT = ROOT / "Higher_order_topology_manuscript_story_report_v3.docx"


def test_story_report_exists_and_is_nontrivial():
    assert REPORT.exists()
    assert REPORT.stat().st_size > 2_000_000


def test_story_report_contains_required_argument_sections():
    doc = Document(REPORT)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    required = [
        "核心科学结论",
        "论文故事线与证据逻辑",
        "结果一：发育过程中出现可重复的高阶共激活重组",
        "结果二：高阶信号与爆发和总体活动水平有关，但不能完全由其解释",
        "结果三：三元组合特异性偏离超过两类低阶约束",
        "结果四：H1拓扑偏离依赖零模型，因此需要限制结论",
        "结果五：效应可跨独立培养复现，并提供有限但可验证的增量信息",
        "统计分析与报告规范",
        "局限性与表述边界",
        "补充分析",
        "文件与复现索引",
    ]
    for heading in required:
        assert heading in text
    assert "PLACEHOLDER" not in text
    assert "TBD" not in text


def test_story_report_embeds_all_main_and_supplementary_figures():
    doc = Document(REPORT)
    assert len(doc.inline_shapes) == 15


def test_story_report_records_core_claim_boundaries():
    doc = Document(REPORT)
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "三元组合特异性" in text
    assert "H1拓扑偏离" in text
    assert "不能解释为直接三神经元突触作用" in text
    assert "143条记录" in text
    assert "P1、P2和P3三次独立培养" in text
