from __future__ import annotations

from release_paths import repository_path

import json
import os
import shutil
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = repository_path('analysis/upgrade')
SOURCE = repository_path('analysis/correction/manuscript/eLife_corrected_clean.docx')
MANUSCRIPT_DIR = ROOT / "manuscript"
REPORT_DIR = ROOT / "report"
FIG_MAIN = ROOT / "figures" / "main" / "Figure_7_higher_order_profiles.png"
FIG_SUPP_1 = ROOT / "figures" / "supplementary" / "Figure_7_figure_supplement_1_O_information.png"
FIG_SUPP_2 = ROOT / "figures" / "supplementary" / "Figure_7_figure_supplement_2_prediction.png"
TABLE_DIR = ROOT / "tables" / "upgrade_v2"

RED = RGBColor(0xFF, 0x00, 0x00)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
TEAL = RGBColor(0x00, 0x78, 0x7A)
LIGHT_BLUE = "D9EAF7"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GREY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, *, bold=False, color=None, size=9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name="Arial", size=10.5, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def mark_paragraph(paragraph, red: bool) -> None:
    if not red:
        return
    for run in paragraph.runs:
        run.font.color.rgb = RED


def replace_paragraph(paragraph, text: str, *, red=False) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    if red:
        run.font.color.rgb = RED


def insert_before(anchor, text: str = "", style: str | None = None, *, red=False):
    p = anchor.insert_paragraph_before(text, style=style)
    mark_paragraph(p, red)
    return p


def insert_after(anchor, text: str = "", style: str | None = None, *, red=False):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    p = anchor._parent.add_paragraph()
    p._p.getparent().remove(p._p)
    new_p.getparent().replace(new_p, p._p)
    if style:
        p.style = style
    if text:
        r = p.add_run(text)
        if red:
            r.font.color.rgb = RED
    return p


def find_paragraph(doc: Document, exact: str):
    for p in doc.paragraphs:
        if p.text.strip() == exact:
            return p
    raise KeyError(exact)


def set_manuscript_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for sty in ["Heading 1", "Heading 2", "Title"]:
        s = doc.styles[sty]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)


def add_display_equation_before(anchor, equation: str, number: str | None = None, *, red=False):
    p = anchor.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = equation if number is None else f"{equation}    ({number})"
    math_run.append(math_text)
    if red:
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "FF0000")
        rpr.append(color)
        math_run.insert(0, rpr)
    math.append(math_run)
    math_para.append(math)
    p._p.append(math_para)
    return p


def add_picture_block(doc: Document, path: Path, caption: str, *, width=6.6, red=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph(caption)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in cp.runs:
        r.italic = True
    mark_paragraph(cp, red)


def build_manuscript(output: Path, *, redline: bool) -> None:
    doc = Document(SOURCE)
    set_manuscript_styles(doc)

    replace_paragraph(
        doc.paragraphs[1],
        "Developmental triplet coactivity is temporally structured but adds limited information beyond pairwise organization in mouse cortical cultures",
        red=redline,
    )

    abstract = (
        "Developing neuronal networks increasingly express coordinated population activity, but simultaneous triplets may arise from firing rates, pairwise dependence, or temporal structure rather than irreducible three-way interactions. We analyzed 143 longitudinal microelectrode-array recordings from three independent mouse cortical preparations between days in vitro (DIV) 14 and 63. Direct triplet support increased most strongly from DIV14 to DIV21. Triplet profiles were temporally continuous and identified individual wells above chance across ages, yet they did not outperform pairwise profiles. Information-theoretic O-information shifted toward redundancy-dominated triplets, but its developmental change was not significant after comparison with either instantaneous count-conditioned pairwise or lag-1 kinetic null ensembles. Similarly, topological and triplet-residual findings depended on the reference model, culture, and sampling specification. Conditional refitting reduced selected upper-tail residual deviations, and synthetic datasets without explicit triplet terms were frequently rejected by the original refitting procedures. A leakage-controlled age-prediction analysis showed a modest full-range improvement when established higher-order features were added to 27 conventional electrophysiological features, whereas O-information added no stable incremental value and later ages remained poorly predicted. These findings show that higher-order representations contain reproducible developmental structure, while most of the measured structure overlaps lower-order organization. Distinguishing direct coactivity, model-relative residuals, and information-theoretic or topological summaries is therefore necessary before interpreting cultured-network data as evidence for irreducible higher-order mechanisms."
    )
    replace_paragraph(doc.paragraphs[9], abstract, red=redline)
    replace_paragraph(
        doc.paragraphs[10],
        "Keywords: cortical culture; neuronal development; microelectrode array; triplet coactivity; O-information; persistent homology; maximum entropy; null model calibration",
        red=redline,
    )
    replace_paragraph(
        doc.paragraphs[12],
        "Longitudinal triplet patterns are reproducible across developmental ages, but most of their information overlaps pairwise structure and does not establish irreducible three-way interactions.",
        red=redline,
    )

    intro_anchor = find_paragraph(doc, "Results")
    old_last_intro = intro_anchor._p.getprevious()
    # Resolve the immediately preceding non-empty paragraph through the high-level list.
    p19 = [p for p in doc.paragraphs if p.text.startswith("Here, we analyzed 143 longitudinal recordings")][0]
    replace_paragraph(
        p19,
        "Here, we analyzed 143 longitudinal recordings from wild-type mouse cortical cultures at eight ages from days in vitro (DIV) 14 to 63. We first characterized direct triplet coactivity and simplicial topology, then compared these measurements with instantaneous pairwise and lag-1 null ensembles preserving the observed population count in every bin. To ask whether alternative higher-order descriptions converged or separated, we added triplet O-information, incidence projections, longitudinal well identification, temporal-continuity tests, and fold-local developmental-age prediction. Recent comparative work in human neuroimaging has shown that higher-order metrics can occupy distinct statistical families and may improve subject identification [50,51]; here we tested whether analogous complementarity is present in longitudinal neuronal recordings rather than assuming that it is. Exact finite-state calculations, conditional refitting, synthetic-data calibration, culture-aware models, and sampling sensitivity were used to separate changes in observed coactivity from conclusions that depend on the adequacy of a statistical reference.",
        red=redline,
    )
    insert_before(
        p19,
        "Information-theoretic measures provide a complementary test because they summarize the balance between redundancy-dominated and synergy-dominated dependencies without choosing a target variable. For three binary variables, O-information is equivalent to interaction information: positive values indicate redundancy-dominated dependence and negative values indicate synergy-dominated dependence [48]. Dynamic variants have been applied to neuronal spiking data [49]. However, O-information is itself a property of the observed joint distribution; its sign does not show that an interaction is causal or irreducible. Comparison with null distributions preserving firing, pairwise moments, population count, and temporal dependence is therefore required before assigning a uniquely higher-order interpretation.",
        red=redline,
    )

    # Replace the previous exploratory prediction result with the leakage-controlled analysis.
    p47 = [p for p in doc.paragraphs if p.text.startswith("As a secondary information check")][0]
    replace_paragraph(
        p47,
        "Because the previously imported 27-feature bridge had undergone pooled preprocessing, we repeated developmental-age prediction from the corrected pre-imputation feature table and restricted every imputation, scaling, and ridge-penalty choice to the training folds. In leave-one-culture-out (LOCO) validation across DIV14–63, mean absolute error (MAE) was 13.685 days for the 27-feature baseline, 12.458 days after adding the established higher-order residual/topological features, 13.480 days after adding O-information features, and 11.987 days after adding both higher-order groups (Fig. 7—figure supplement 2A,D). Relative to the baseline, the whole-well bootstrap MAE improvements were 1.227 days (95% CI 0.052–2.513) for the established higher-order features and 1.698 days (95% CI 0.496–2.912) for all higher-order features, but only 0.205 days (95% CI −0.567 to 0.963) for O-information alone (Fig. 7—figure supplement 2C). O-information did not improve the established higher-order model reliably (increment 0.471 days, 95% CI −0.314 to 1.265). All DIV21–63 LOCO models had negative pooled R², so no claim of later-age prediction was retained (Fig. 7—figure supplement 2B).",
        red=redline,
    )

    discussion_anchor = find_paragraph(doc, "Discussion")
    insert_before(
        discussion_anchor,
        "Higher-order profiles are temporally structured but add limited information beyond pairwise coactivity",
        style="Heading 2",
        red=redline,
    )
    results_paragraphs = [
        "Higher-order summaries may change with age yet still describe the same underlying organization as pairwise coactivity. We therefore compared pairwise edges, direct triplet support, null-relative residuals, and O-information in a common incidence-projection framework (Fig. 7A). Across recordings, the median Spearman correlation between pairwise edge weight and the corresponding direct-triplet edge projection was 0.979, indicating near-overlap at this descriptive level (Fig. 7B). In contrast, pairwise weight correlated negatively with signed residual projection (median ρ = −0.448), while signed residual and O-information-deviation projections were closely related (median ρ = 0.825). The metric taxonomy thus separated direct coactivity, which closely tracked pairwise support, from null-relative representations, which emphasized how triplet allocation departed from lower-order expectations.",
        "To determine whether these representations preserved developmental order within individual networks, we used the 17 wells with complete longitudinal measurements and compared adjacent-age similarities with 10,000 within-well age-order permutations. Observed adjacent-age Spearman correlations exceeded the permuted expectation for pairwise edges (0.471 versus 0.374, P = 1.0 × 10⁻⁴), direct triplets (0.454 versus 0.359, P = 1.0 × 10⁻⁴), O-information (0.349 versus 0.241, P = 1.0 × 10⁻⁴), signed residuals (0.199 versus 0.102, P = 1.0 × 10⁻⁴), absolute residuals (0.178 versus 0.135, P = 0.0031), and O-information deviations (0.218 versus 0.126, P = 1.0 × 10⁻⁴; Fig. 7C). Thus, even null-relative profiles retained continuity across successive ages rather than fluctuating as arbitrary triplet rankings.",
        "We next asked whether each representation contained a stable well-specific signature across ages. Each recording was matched only against recordings from the same preparation but a different age, excluding the query age; the analytic chance level was 0.176. Pairwise, direct-triplet, and O-information profiles identified the correct well above chance, with top-1 accuracies of 0.747 (95% CI 0.705–0.791), 0.730 (0.685–0.775), and 0.593 (0.546–0.640), respectively (all permutation P = 1.0 × 10⁻⁴; Fig. 7D). Direct-triplet accuracy was nevertheless 0.0168 lower than the pairwise value (95% CI −0.0326 to −0.0011), and concatenating pairwise and triplet projections produced no reliable improvement (difference −0.0032, 95% CI −0.0074 to 0.0011; Fig. 7E). Higher-order profiles therefore preserved recording-specific structure, but this structure was not more discriminative than pairwise support.",
        "O-information was used to distinguish redundancy-dominated from synergy-dominated triplets. Mean observed O-information increased from 0.0091 bits at DIV14 to 0.0734 bits at DIV21 (adjusted P = 0.00179), and the fraction of redundancy-dominated triplets increased by 0.621 (adjusted P = 6.52 × 10⁻³⁰; Fig. 7—figure supplement 1A,B). The corresponding early change in observed-minus-null mean O-information was not significant for the instantaneous reference (0.00070 bits, adjusted P = 0.405) or the kinetic reference (0.00094 bits, adjusted P = 0.598; Fig. 7F and Fig. 7—figure supplement 1C,D); P3 also showed a negative median early change under both references. All preregistered O-information claim gates were therefore false. The observed redundancy shift is a valid distributional description, but the tested lower-order constraints were sufficient to reproduce its developmental change.",
    ]
    for text in results_paragraphs:
        insert_before(discussion_anchor, text, red=redline)

    # Add interpretive discussion before Limitations.
    limitations_anchor = find_paragraph(doc, "Limitations")
    discussion_additions = [
        "The profile analyses clarify what the developmental signal represents. Direct triplet support was almost collinear with its pairwise projection, so its excellent longitudinal continuity and above-chance well identification are not evidence that three-way statistics supply an independent organizational axis. This differs from reports in human fMRI in which some higher-order methods improved individual fingerprinting [50,51]. The discrepancy is informative rather than contradictory: the spatial scale, recording modality, null construction, and target task differ substantially, and comparative benchmarking is needed precisely because no higher-order metric is guaranteed to add information in every system.",
        "Null-relative metrics occupied a more distinct position in the metric taxonomy and retained temporal continuity, but distinctness did not translate into superior identification. This separates three questions that are often conflated: whether a representation differs mathematically from a pairwise graph, whether it is stable within a biological preparation, and whether it adds predictive or discriminative information. Our data answer the first two questions positively for selected residual representations, but provide no general support for the third. The resulting contribution is a boundary on interpretation: measurable triplet structure can be reproducible while remaining largely explainable by lower-order organization.",
        "The O-information analysis provides an independent information-theoretic view of the same boundary. Positive O-information indicates that a triplet's joint distribution is redundancy-dominated, not that a dedicated triadic coupling generated it [48]. Because both count-conditioned references reproduced the early change in mean O-information, the raw developmental shift can arise from the increasing firing, pairwise coordination, population recruitment, and short-lag structure already encoded by those references. The result does not imply that synergy is absent from cultured networks; it shows that this recording duration, bin width, node selection, and estimator do not provide robust evidence for an age-dependent increase in residual redundancy or synergy beyond the tested constraints.",
        "Leakage-controlled prediction led to the same qualified conclusion. Established residual and topological features modestly reduced full-range LOCO error, but much of this gain was associated with the large DIV14-to-DIV21 transition, and all later-age models had negative held-out-culture R². O-information neither improved the conventional baseline reliably nor added stable information beyond the established higher-order group. These observations argue against describing the metrics as general biomarkers of maturation. A stronger test would compare preparations or perturbations with similar firing and pairwise organization but different developmental outcomes, using an external cohort and a fully prospective feature pipeline.",
    ]
    for text in discussion_additions:
        insert_before(limitations_anchor, text, red=redline)

    # Update prediction discussion to remove the superseded globally preprocessed analysis.
    pred_disc = [p for p in doc.paragraphs if p.text.startswith("The secondary age-prediction analysis provides")][0]
    replace_paragraph(
        pred_disc,
        "The developmental-age analysis is therefore best treated as a secondary information test rather than a validation endpoint. Rebuilding it from the pre-imputation table removed the clearest avoidable preprocessing leakage, but it did not create reliable resolution of post-DIV21 age. The modest full-range improvement from established higher-order features and the absence of stable incremental O-information are consistent with the profile analyses: higher-order summaries can carry reproducible signal without defining a broadly superior maturation coordinate.",
        red=redline,
    )

    last_limit = [p for p in doc.paragraphs if p.text.startswith("The null families are separate and not nested")][0]
    replace_paragraph(
        last_limit,
        "The null families are separate and not nested; neither jointly constrains instantaneous and lag-1 moments. Both use finite ensembles of 100 realizations. Target-moment RMSE below 0.02 is an operational diagnostic, not proof that mismatch cannot explain a residual. The exact/refit benchmark includes nine selected recordings, and the synthetic calibration is restricted to six frozen prototypes. O-information estimates from sparse binary triplets are sensitive to finite-sample probability estimation and do not constitute a partial-information decomposition or a causal interaction estimate. Longitudinal identification was restricted to 17 complete wells, compared candidates only within preparation, and excluded the query age; the fixed-12 residual-continuity analysis also uses electrodes selected for completeness across the trajectory. The leakage-controlled prediction began from a pre-imputation table, but some imported conventional features still inherited upstream DIV-dependent availability masks. No independent cohort or intervention was used to test biological specificity.",
        red=redline,
    )

    conclusion = [p for p in doc.paragraphs if p.text.startswith("Developing cortical cultures showed an early expansion")][0]
    replace_paragraph(
        conclusion,
        "Developing cortical cultures showed an early expansion of direct triplet coactivity and a shift toward redundancy-dominated joint activity. These higher-order profiles were temporally continuous and retained well-specific information across ages, but direct triplet organization closely overlapped pairwise structure and did not improve longitudinal identification. Developmental changes in O-information, H₁ persistence, and standardized residuals were sensitive to lower-order reference models, culture, or sampling specification, and O-information added no stable predictive value beyond established features. The study therefore supports reproducible higher-order representations of development while limiting their interpretation as irreducible three-way mechanisms. Pairwise overlap, null-relative deviation, and causal interaction are distinct claims and should be tested separately.",
        red=redline,
    )

    # Update and extend Methods.
    baseline_method = [p for p in doc.paragraphs if p.text.startswith("The secondary prediction comparator comprised")][0]
    replace_paragraph(
        baseline_method,
        "The developmental-age comparator comprised the original 27 well-level electrophysiological features: firing and network-burst measures, STTC connectivity, topology and node roles, NMF component count, covariance effective rank, and finite-horizon control energy [32,40–45]. For V2, these columns were taken from the corrected filtered table before median imputation. The feature matrix contained 84 missing cells across 142 retained recordings. Within every outer and inner validation split, missing values were imputed with training-fold medians and all columns were standardized using training-fold means and standard deviations. Some source features retained upstream DIV-dependent availability masks; this residual limitation is reported explicitly."
        , red=redline)

    robustness_method = [p for p in doc.paragraphs if p.text.startswith("Ridge regression predicted DIV")][0]
    replace_paragraph(
        robustness_method,
        "Ridge regression predicted DIV under nested leave-one-culture-out (LOCO) validation, with leave-one-well-out analysis retained as sensitivity analysis [36,39]. Four matched feature sets were compared: the 27-feature baseline; baseline plus established higher-order residual and topological features; baseline plus O-information features; and baseline plus both groups. Inner grouped folds selected among 13 penalties from 10⁻³ to 10³. Imputation, scaling, and tuning were repeated independently inside each fold. Full-range (DIV14–63) and late-range (DIV21–63) models were trained separately. MAE improvement was comparator MAE minus augmented-model MAE; 95% percentile intervals used 5,000 culture-stratified bootstrap samples of complete wells. A positive information claim required a positive lower confidence bound and non-negative held-out R² in the target age range.",
        red=redline,
    )

    stats_anchor = find_paragraph(doc, "Statistical analysis")
    methods = [
        ("Triplet incidence projection and longitudinal identification", "heading"),
        ("To compare representations defined on nodes, edges, and triplets, each triplet-valued vector was projected to its incident edges or nodes. The value assigned to an edge was the arithmetic mean over all retained triplets containing that edge; node projections were defined analogously over all incident triplets. Direct-support, signed-residual, absolute-residual, O-information, and observed-minus-null O-information projections were computed without thresholding the source vectors. Cross-metric taxonomy used Spearman correlations within recording and summarized each pair by the median across recordings.", "normal"),
        ("Longitudinal identification used the 17 wells with complete measurements at all eight ages. For each query, candidate recordings were restricted to the same culture, the same age was excluded, and similarity was Spearman correlation between feature vectors. The predicted well was the candidate with greatest similarity. We recorded top-1 accuracy, true-match rank, and the true-minus-best-impostor similarity margin. Permutation P values used 10,000 label permutations and the (b + 1)/(B + 1) correction [46]. Confidence intervals and pairwise differences between representations used 5,000 culture-stratified whole-well bootstrap samples. Temporal continuity used the mean similarity of adjacent ages within wells and 10,000 within-well permutations of age order.", "normal"),
        ("Third-order O-information", "heading"),
        ("For each dynamic 16-electrode recording, spike events were converted to binary variables in the primary 25-ms bins. For every triplet (Xi, Xj, Xk), the eight joint-state probabilities were calculated from observed first-, second-, and third-order moments, and Shannon entropies used base-2 logarithms. Third-order O-information was calculated as", "normal"),
        ("EQUATION_OINFO", "equation"),
        ("where H denotes Shannon entropy. Positive values indicate redundancy-dominated statistical dependence and negative values indicate synergy-dominated dependence [48]. For clarity, this sign classification does not identify a causal or irreducible interaction. Joint probabilities were reconstructed as p111 = pijk, p110 = pij − pijk, p101 = pik − pijk, p011 = pjk − pijk, p100 = pi − pij − pik + pijk, p010 = pj − pij − pjk + pijk, p001 = pk − pik − pjk + pijk, and p000 = 1 − pi − pj − pk + pij + pik + pjk − pijk. Probabilities were clipped only within floating-point tolerance and renormalized.", "normal"),
        ("Each observed recording was compared with 100 instantaneous count-conditioned pairwise simulations and 100 lag-1 kinetic count-conditioned simulations generated with the same primary specifications as the residual analysis. Recording-level outcomes were mean O-information, positive and negative O-information mass, redundancy-dominated and synergy-dominated triplet fractions, and observed-minus-null deviations. Culture-aware longitudinal models and explicit P1–P3 early-change summaries were applied to the deviations. A developmental O-information claim required an adjusted early contrast below 0.05, concordant early-change direction in all three cultures, and agreement under both null families.", "normal"),
        ("V2 decision rules and sample sizes", "heading"),
        ("Full-array direct node/edge/triplet profiles used all 143 recordings. Dynamic 16-node O-information used 142 recordings because one recording did not meet the pre-existing 16-electrode eligibility rule. Fixed 12-node residual and temporal-continuity analyses used 143 recordings, while cross-age identification and continuity statistics were restricted to 17 complete wells. The full-range prediction used 142 recordings and the late-range analysis used 124. Missing outcomes were not replaced by zeros. Unless otherwise stated, all V2 null analyses used 100 realizations, permutation tests used 10,000 permutations, and bootstrap intervals used 5,000 culture-stratified whole-well resamples.", "normal"),
    ]
    for text, kind in methods:
        if kind == "heading":
            insert_before(stats_anchor, text, style="Heading 2", red=redline)
        elif kind == "equation":
            add_display_equation_before(
                stats_anchor,
                "Ω(Xi, Xj, Xk) = H(Xi) + H(Xj) + H(Xk) + H(Xi, Xj, Xk)",
                red=redline,
            )
            add_display_equation_before(
                stats_anchor,
                "− H(Xi, Xj) − H(Xi, Xk) − H(Xj, Xk)",
                red=redline,
            )
        else:
            insert_before(stats_anchor, text, red=redline)

    stats_text = [p for p in doc.paragraphs if p.text.startswith("Longitudinal outcomes were analyzed")][0]
    replace_paragraph(
        stats_text,
        stats_text.text + " V2 longitudinal-identification and continuity P values were permutation-based, with confidence intervals obtained by culture-stratified whole-well bootstrap. O-information early contrasts were evaluated with the same culture-aware mixed-effects framework and Benjamini–Hochberg adjustment as the primary outcomes. Prediction comparisons used paired whole-well bootstrap distributions from identical held-out observations. Statistical significance was not used alone to pass a claim gate when cross-culture direction, null-family agreement, or held-out predictive performance failed the prespecified criterion.",
        red=redline,
    )

    # Fill the previously empty abbreviation section.
    decl_anchor = find_paragraph(doc, "Declarations")
    insert_before(
        decl_anchor,
        "DIV, days in vitro; H₁, first homology group; HO, higher order; LOCO, leave one culture out; MAE, mean absolute error; MEA, microelectrode array; NMF, non-negative matrix factorization; O-information, organizational information; PCA, principal component analysis; RMSE, root mean square error; STTC, spike-time tiling coefficient.",
        red=redline,
    )

    # Add verified references before the figure section.
    ref_anchor = find_paragraph(doc, "Figure 1")
    refs = [
        "48. Rosas FE, Mediano PAM, Gastpar M, Jensen HJ. Quantifying high-order interdependencies via multivariate extensions of the mutual information. Phys Rev E. 2019;100:032305. doi:10.1103/PhysRevE.100.032305.",
        "49. Stramaglia S, Scagliarini T, Daniels BC, Marinazzo D. Quantifying dynamical high-order interdependencies from the O-information: an application to neural spiking dynamics. Front Physiol. 2021;11:595736. doi:10.3389/fphys.2020.595736.",
        "50. Santoro A, Battiston F, Lucas M, et al. Higher-order connectomics of human brain function reveals local topological signatures of task decoding, individual identification, and behavior. Nat Commun. 2024;15:10244. doi:10.1038/s41467-024-54472-y.",
        "51. Santoro A, Neri M, Poetto S, et al. Charting higher-order models of brain function beyond pairwise interactions. Nat Commun. 2026;17:9207. doi:10.1038/s41467-026-75959-w.",
    ]
    for ref in refs:
        insert_before(ref_anchor, ref, red=redline)

    # Add new eLife-style figure legends and figures at the end.
    doc.add_heading("Figure 7", level=2)
    p = doc.add_paragraph(
        "Higher-order profiles are temporally structured but largely overlap pairwise information. (A) Analysis framework linking node, edge, and triplet representations through incidence projections, within-well temporal continuity, and cross-age well identification. (B) Median recording-wise Spearman correlations among pairwise, direct-triplet, residual, and O-information projections. (C) Observed adjacent-age similarity compared with within-well age-order permutations for 17 complete wells. (D) Top-1 cross-age well-identification accuracy; the dashed line denotes the analytic chance level of 0.176. Error bars are 95% culture-stratified whole-well bootstrap confidence intervals. (E) Accuracy difference from the pairwise baseline; intervals crossing zero indicate no reliable improvement. (F) Change from DIV14 to DIV21 in mean observed-minus-null O-information under the instantaneous and kinetic references, shown overall and by culture."
    )
    mark_paragraph(p, redline)
    add_picture_block(doc, FIG_MAIN, "Figure 7. Higher-order profile comparison.", red=redline)

    doc.add_heading("Figure 7—figure supplement 1", level=2)
    p = doc.add_paragraph(
        "Observed and null-relative O-information across development. (A) Mean observed triplet O-information by DIV. (B) Fraction of triplets with positive O-information. (C,D) Recording-level observed-minus-null mean O-information under instantaneous count-conditioned pairwise (C) and lag-1 kinetic count-conditioned (D) references. Points denote wells and colors denote independent culture preparations; boxes summarize the distributions."
    )
    mark_paragraph(p, redline)
    add_picture_block(doc, FIG_SUPP_1, "Figure 7—figure supplement 1. O-information summaries.", red=redline)

    doc.add_heading("Figure 7—figure supplement 2", level=2)
    p = doc.add_paragraph(
        "Leakage-controlled developmental-age prediction. (A,B) LOCO MAE for the 27-feature baseline and three augmented feature sets across all ages (A) and DIV21–63 (B). (C) Whole-well bootstrap MAE improvement relative to the 27-feature baseline for the full age range; points are observed improvements and lines are 95% confidence intervals. (D) Held-out DIV predictions from the all-higher-order model, colored by culture; the dashed line is identity. Negative late-range R² values precluded a later-age prediction claim."
    )
    mark_paragraph(p, redline)
    add_picture_block(doc, FIG_SUPP_2, "Figure 7—figure supplement 2. Developmental-age prediction.", red=redline)

    # Normalize fonts without overriding redlining.
    for p in doc.paragraphs:
        for r in p.runs:
            if p.style.name == "Title":
                set_run_font(r, "Arial", 18, color=r.font.color.rgb, bold=True)
            elif p.style.name.startswith("Heading"):
                set_run_font(r, "Arial", 11.5 if p.style.name == "Heading 2" else 13, color=r.font.color.rgb, bold=True)
            else:
                set_run_font(r, "Arial", 10.5, color=r.font.color.rgb)

    doc.core_properties.title = doc.paragraphs[1].text
    doc.core_properties.subject = "eLife manuscript upgraded with comparative higher-order analyses"
    doc.save(output)


def configure_report_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for sty, size, color in [("Title", 20, BLUE), ("Heading 1", 15, BLUE), ("Heading 2", 12, TEAL), ("Heading 3", 10.5, TEAL)]:
        s = doc.styles[sty]
        s.font.name = "Arial"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.color.rgb = color
        s.font.bold = True


def report_table(doc: Document, headers: list[str], rows: Iterable[Iterable], widths=None, note: str | None = None):
    rows = list(rows)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[j], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9)
        set_cell_shading(t.rows[0].cells[j], "1F4E79")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value, size=8.5)
            if i % 2:
                set_cell_shading(cells[j], "F7FAFC")
    if note:
        p = doc.add_paragraph(note)
        for r in p.runs:
            r.italic = True
            r.font.size = Pt(8.5)
    return t


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def fmt(x, digits=3):
    try:
        x = float(x)
        if abs(x) < 0.001 and x != 0:
            return f"{x:.2e}"
        return f"{x:.{digits}f}"
    except Exception:
        return str(x)


def build_report(output: Path) -> None:
    doc = Document()
    configure_report_styles(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("培养小鼠皮层网络高阶拓扑研究\nV2升级分析完整报告")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Correction_elife_V2 | 2026-09-13").italic = True
    doc.add_paragraph()
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_shading(box.cell(0, 0), LIGHT_BLUE)
    set_cell_text(
        box.cell(0, 0),
        "一句话结论：培养网络的高阶表征具有真实、可重复的纵向结构，但直接三元共活动与成对结构高度重叠；原始O-information的冗余化趋势可由成对、群体计数和短时序约束再现，因此目前证据支持“可测量的高阶表征”，不支持“发育过程中增强的不可约三元机制”。",
        bold=True,
        size=11,
    )

    doc.add_heading("1. 本次升级解决的核心问题", level=1)
    for x in [
        "高阶指标是否只是把成对关系换一种形式表达，还是形成可区分的统计表征？",
        "高阶模式是否在同一培养孔的长期发育中保持连续，并能跨年龄识别该孔？",
        "O-information所显示的冗余/协同倾向，是否超过成对关系、每个时间窗的群体放电数以及一阶时间依赖所能解释的范围？",
        "高阶指标是否在严格避免折外信息泄漏的条件下，提高发育年龄预测？",
    ]:
        add_number(doc, x)
    doc.add_paragraph(
        "这四个问题把论文从“发现若干复杂指标随DIV变化”提升为“比较不同表征、检验其独立信息，并给出清晰的推断边界”。这种升级与Santoro等对高阶指标进行分类和功能比较的思路一致，但分析对象、时间尺度和科学问题均不同：对方研究人脑fMRI个体指纹与行为，本研究检验培养神经网络发育中的电极级三元共活动。"
    )

    doc.add_heading("2. 数据合同与完整性", level=1)
    report_table(
        doc,
        ["层级", "样本量", "用途", "处理原则"],
        [
            ["全64电极", "143条记录", "直接放电、成对和三元支持投影", "不补零，不删除低值三元"],
            ["动态16电极", "142条记录", "O-information及双零模型", "沿用既定电极资格；1条不满足规则"],
            ["固定12电极", "143条记录", "完整三元向量、残差和时序连续性", "全部220个三元，不截取top-400"],
            ["完整纵向孔", "17个孔", "跨年龄识别与连续性", "同培养内匹配，排除同年龄候选"],
            ["预测", "142条；晚期124条", "LOCO/LOWO年龄预测", "折内插补、标准化和调参"],
        ],
        note="生物学重复为3次独立培养P1、P2、P3；143条记录不是143个独立培养。"
    )
    add_bullet(doc, "记录时长为299.75–300 s，主分析使用25 ms时间窗，不对短记录进行填充。")
    add_bullet(doc, "所有V2主零模型均使用100次模拟；每条记录保存随机种子、模型参数和压缩检查点。")
    add_bullet(doc, "原始输入建立SHA256清单；V2仅写入新目录，不回写纠正后的源项目。")

    doc.add_heading("3. 新增方法", level=1)
    doc.add_heading("3.1 三元关联投影与指标分类", level=2)
    doc.add_paragraph(
        "三元指标与成对指标的维度不同，不能直接逐元素比较。对每个三元向量，分别投影到其包含的边和节点：某条边的投影值等于所有包含该边的三元值的平均值；节点投影同理。随后在每条记录内计算Spearman相关，跨记录取中位数，形成指标分类矩阵。该方法回答“两个指标是否强调相同的空间组织”，而不是比较它们的数值单位。"
    )
    add_display_equation_before(doc.add_paragraph(), "P_e(x) = [1 / |T(e)|] Σ_{t∈T(e)} x_t")

    doc.add_heading("3.2 纵向识别与时序连续性", level=2)
    doc.add_paragraph(
        "对17个八个年龄均完整的培养孔，将某一年龄的记录作为查询。候选集只包含同一次独立培养、不同年龄的其他记录，避免培养批次差异把识别任务变得过于容易，并排除同年龄候选。以Spearman相关最高者作为预测孔，统计top-1准确率、真匹配排名和真匹配相对最佳错误匹配的相似度优势。准确率的理论随机水平为1/5.667=0.176。显著性使用10,000次标签置换；置信区间按培养分层、整孔重采样5,000次。"
    )
    doc.add_paragraph(
        "时序连续性比较真实相邻DIV的平均相似度与同一孔内随机打乱年龄顺序后的期望。该检验并不要求指标单调变化，只检验相邻发育阶段是否比随机年龄排列更相似。"
    )

    doc.add_heading("3.3 三阶O-information", level=2)
    doc.add_paragraph(
        "对三个二元电极变量计算三阶O-information。它是对称的，不指定预测变量或靶变量；正值表示冗余占主导，负值表示协同占主导。这里的“协同/冗余”属于联合概率分布的统计描述，不等于细胞之间存在相应方向的因果连接。"
    )
    p = doc.add_paragraph()
    add_display_equation_before(
        p,
        "Ω(Xi, Xj, Xk) = H(Xi) + H(Xj) + H(Xk) + H(Xi, Xj, Xk)",
    )
    add_display_equation_before(p, "− H(Xi, Xj) − H(Xi, Xk) − H(Xj, Xk)")
    doc.add_paragraph(
        "为判断原始O-information趋势是否具有独立的高阶含义，每条记录分别与100个瞬时计数条件化成对最大熵模拟和100个lag-1动力学计数条件化模拟比较。只有早期差异显著、P1/P2/P3方向一致且两个零模型均支持时，才通过主张门槛。"
    )

    doc.add_heading("3.4 无泄漏年龄预测", level=2)
    doc.add_paragraph(
        "从纠正后、尚未插补的27特征表出发，在每个外层和内层训练折内独立执行中位数插补、z标准化与岭回归参数选择。比较四组特征：传统27项、传统+原高阶特征、传统+O-information、传统+全部高阶特征。主验证为leave-one-culture-out（LOCO）；晚期DIV21–63单独训练，防止DIV14与其余阶段的巨大差异掩盖晚期性能。"
    )

    doc.add_heading("4. 主要结果", level=1)
    doc.add_heading("4.1 直接三元共活动与成对结构几乎同构", level=2)
    corr = pd.read_csv(TABLE_DIR / "v2_cross_metric_median_correlation_matrix.csv", index_col=0)
    report_table(
        doc,
        ["指标对", "跨记录中位Spearman ρ", "解释"],
        [
            ["成对边 vs 直接三元边投影", "0.979", "直接三元空间分布几乎由成对支持排序概括"],
            ["成对边 vs 有符号残差投影", "−0.448", "残差强调与强连接不同的三元配置"],
            ["有符号残差 vs O-info偏差投影", "0.825", "两种零模型相对量形成相近的独立指标族"],
            ["成对边 vs O-info投影", "0.450", "原始信息量与成对结构部分重叠"],
        ],
    )
    doc.add_paragraph(
        "最重要的不是‘相关很高’本身，而是它改变了文章的表述：直接三元共活动虽然是高阶对象，但其空间组织没有形成独立于成对网络的新轴；真正与成对结构分离的是扣除零模型后的相对量。"
    )

    doc.add_heading("4.2 所有表征都保留发育连续性", level=2)
    report_table(
        doc,
        ["表征", "真实相邻ρ", "置换期望ρ", "P"],
        [
            ["成对边", "0.471", "0.374", "1.0×10⁻⁴"],
            ["直接三元", "0.454", "0.359", "1.0×10⁻⁴"],
            ["O-information", "0.349", "0.241", "1.0×10⁻⁴"],
            ["有符号残差", "0.199", "0.102", "1.0×10⁻⁴"],
            ["绝对残差", "0.178", "0.135", "0.0031"],
            ["O-info偏差", "0.218", "0.126", "1.0×10⁻⁴"],
        ],
        note="17个完整纵向孔；10,000次孔内年龄顺序置换。"
    )
    doc.add_paragraph(
        "该结果是V2中最明确的正向结论：即使残差类指标没有优于成对网络，它们也不是随机噪声，仍然沿真实发育顺序连续变化。"
    )

    doc.add_heading("4.3 高阶表征可以识别培养孔，但不优于成对网络", level=2)
    report_table(
        doc,
        ["表征", "Top-1准确率", "95% CI", "置换P", "相对成对差值"],
        [
            ["放电节点", "0.745", "0.701–0.789", "1.0×10⁻⁴", "—"],
            ["成对边", "0.747", "0.705–0.791", "1.0×10⁻⁴", "0（基准）"],
            ["直接三元", "0.730", "0.685–0.775", "1.0×10⁻⁴", "−0.0168 [−0.0326, −0.0011]"],
            ["O-information", "0.593", "0.546–0.640", "1.0×10⁻⁴", "−0.1534 [−0.1975, −0.1134]"],
            ["成对+直接三元", "0.744", "—", "—", "−0.0032 [−0.0074, 0.0011]"],
            ["成对+O-info", "0.696", "—", "—", "−0.0504 [−0.0746, −0.0231]"],
        ],
        note="候选仅来自同一独立培养且年龄不同；理论机会水平0.176。"
    )
    doc.add_paragraph(
        "因此不能写成‘高阶方法提高了培养孔指纹识别’。更准确的结论是：三元和O-information都包含稳定个体信息，但该信息大多已经存在于成对结构中。"
    )

    doc.add_heading("4.4 原始O-information显著冗余化，但零模型相对变化不显著", level=2)
    report_table(
        doc,
        ["结果", "DIV14", "DIV21/早期变化", "调整P", "是否通过主张门槛"],
        [
            ["平均原始O-information（bit）", "0.0091", "0.0734；Δ=0.0642", "0.00179", "否：零模型可再现"],
            ["冗余占主导三元比例", "—", "Δ=0.621", "6.52×10⁻³⁰", "否：属于原始分布"],
            ["瞬时零模型偏差（bit）", "—", "Δ=0.00070", "0.405", "否"],
            ["动力学零模型偏差（bit）", "—", "Δ=0.00094", "0.598", "否"],
        ],
    )
    doc.add_paragraph(
        "这不是‘没有结果’，而是强约束结论：随着网络成熟，观察到的联合活动确实更偏向冗余，但这一变化不需要额外的不可约三元项即可解释。P3在两个偏差指标上的早期中位变化均为负，进一步阻止了跨培养一致性主张。"
    )

    doc.add_heading("4.5 预测只在全龄程获得有限增益，晚期预测失败", level=2)
    report_table(
        doc,
        ["LOCO模型", "全龄程MAE（天）", "全龄程R²", "晚期MAE（天）", "晚期R²"],
        [
            ["传统27特征", "13.685", "−0.048", "22.584", "−22.678"],
            ["+ 原高阶特征", "12.458", "0.141", "16.938", "−6.064"],
            ["+ O-information", "13.480", "0.019", "40.154", "−134.171"],
            ["+ 全部高阶", "11.987", "0.154", "22.756", "−22.286"],
        ],
        note="晚期所有模型R²均为负，因此晚期预测主张未通过。"
    )
    report_table(
        doc,
        ["全龄程MAE改进", "点估计（天）", "95%整孔bootstrap CI", "解释"],
        [
            ["传统 → 原高阶", "1.227", "0.052–2.513", "有限但稳定"],
            ["传统 → O-info", "0.205", "−0.567–0.963", "不稳定"],
            ["传统 → 全部高阶", "1.698", "0.496–2.912", "有限但稳定"],
            ["原高阶 → 全部高阶", "0.471", "−0.314–1.265", "O-info无可靠增量"],
        ],
    )

    doc.add_heading("5. 论文主张门槛", level=1)
    report_table(
        doc,
        ["候选主张", "结果", "论文可用表述"],
        [
            ["高阶指标包含纵向结构", "通过", "多种高阶/残差表征的真实相邻年龄相似度高于孔内年龄置换"],
            ["高阶指标包含孔特异信息", "通过", "直接三元和O-info跨年龄识别率显著高于机会水平"],
            ["高阶指标优于成对指纹", "不通过", "直接三元略低于成对；拼接不增加准确率"],
            ["发育导致超出成对约束的O-info变化", "不通过", "原始冗余增加，但两个零模型偏差均不显著且培养方向不一致"],
            ["O-info提高年龄预测", "不通过", "O-info单独及相对原高阶特征均无稳定增量"],
            ["高阶特征可靠解析DIV21–63", "不通过", "所有晚期LOCO模型R²为负"],
        ],
    )

    doc.add_heading("6. 图件及论文故事线", level=1)
    doc.add_paragraph(
        "新增Figure 7被设计成论文的收束图，而不是再增加一组独立现象。A先建立统一投影与验证框架；B回答不同指标是否同类；C证明其具有发育连续性；D–E检验是否提供超越成对关系的指纹信息；F用双零模型限制O-information的机制解释。两个附图分别展开O-information全龄程分布与无泄漏预测。"
    )
    add_picture_block(doc, FIG_MAIN, "图R1. 新增主图Figure 7。", width=6.7)
    add_picture_block(doc, FIG_SUPP_1, "图R2. Figure 7—figure supplement 1：O-information。", width=6.7)
    add_picture_block(doc, FIG_SUPP_2, "图R3. Figure 7—figure supplement 2：预测。", width=6.7)

    doc.add_heading("7. 与V1相比，论文结论如何变化", level=1)
    report_table(
        doc,
        ["V1倾向", "V2修正", "科学意义"],
        [
            ["主要围绕三元支持、残差和H₁的发育变化", "加入表征分类、纵向指纹、O-info及无泄漏预测", "从指标罗列转向独立信息检验"],
            ["高阶残差可能被理解为独立机制", "明确区分直接高阶对象、零模型偏差和不可约因果作用", "避免概念跨越"],
            ["预测桥存在全局预处理限制", "从插补前表重建折内预处理", "消除主要可避免的信息泄漏"],
            ["缺少与成对结构的正面对照", "直接三元投影与成对ρ=0.979，识别不优于成对", "给出诚实、可检验的限制性结论"],
        ],
    )

    doc.add_heading("8. 局限与下一步", level=1)
    for x in [
        "MEA电极记录的是多单位群体活动；三元电极统计不能直接映射为三个单神经元之间的突触机制。",
        "仅有3次独立培养，虽然纵向孔数为17–18，但跨培养推广仍有限。",
        "O-information在稀疏二元变量中受有限样本影响；它不是完整的partial information decomposition。",
        "两个零模型分别约束瞬时和lag-1统计，但不是联合嵌套模型，也不能穷尽共同输入和更长记忆。",
        "预测仍存在部分源特征的DIV依赖可用性掩码，且没有外部培养批次或扰动队列。",
        "如果未来增加敲除或药理数据，最有价值的设计不是再证明放电率改变，而是匹配传统活动/成对指标后检验高阶残差是否仍能区分条件。",
    ]:
        add_bullet(doc, x)

    doc.add_heading("9. 交付物索引", level=1)
    report_table(
        doc,
        ["类型", "位置"],
        [
            ["清洁版论文", str(MANUSCRIPT_DIR / "eLife_V2_clean.docx")],
            ["红字修改版", str(MANUSCRIPT_DIR / "eLife_V2_redmarked.docx")],
            ["主图与单面板", str(ROOT / "figures")],
            ["源数据/统计工作簿", str(ROOT / "spreadsheets")],
            ["全部CSV结果", str(TABLE_DIR)],
            ["逐记录检查点", str(ROOT / "intermediate")],
            ["代码与测试", str(ROOT / "code") + "；" + str(ROOT / "tests")],
            ["输入哈希与QA", str(ROOT / "audit") + "；" + str(ROOT / "qa")],
        ],
    )

    doc.add_heading("10. 参考方法文献", level=1)
    refs = [
        "Rosas FE, Mediano PAM, Gastpar M, Jensen HJ. Quantifying high-order interdependencies via multivariate extensions of the mutual information. Physical Review E. 2019;100:032305. doi:10.1103/PhysRevE.100.032305.",
        "Stramaglia S, Scagliarini T, Daniels BC, Marinazzo D. Quantifying dynamical high-order interdependencies from the O-information: an application to neural spiking dynamics. Frontiers in Physiology. 2021;11:595736. doi:10.3389/fphys.2020.595736.",
        "Santoro A, Battiston F, Lucas M, et al. Higher-order connectomics of human brain function reveals local topological signatures of task decoding, individual identification, and behavior. Nature Communications. 2024;15:10244. doi:10.1038/s41467-024-54472-y.",
        "Santoro A, Neri M, Poetto S, et al. Charting higher-order models of brain function beyond pairwise interactions. Nature Communications. 2026;17:9207. doi:10.1038/s41467-026-75959-w.",
    ]
    for x in refs:
        p = doc.add_paragraph(x)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.size = Pt(8.5)

    # Footer and pagination fields.
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Correction_elife_V2 | analysis report | ")
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        footer._p.append(fld)

    doc.core_properties.title = "培养小鼠皮层网络高阶拓扑研究V2升级分析完整报告"
    doc.core_properties.subject = "eLife higher-order topology upgrade"
    doc.core_properties.author = "OpenAI Codex (analysis and document assistance)"
    doc.save(output)


def main():
    MANUSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    build_manuscript(MANUSCRIPT_DIR / "eLife_V2_clean.docx", redline=False)
    build_manuscript(MANUSCRIPT_DIR / "eLife_V2_redmarked.docx", redline=True)
    build_report(REPORT_DIR / "高阶拓扑论文V2升级完整报告.docx")
    print(json.dumps({
        "clean": str(MANUSCRIPT_DIR / "eLife_V2_clean.docx"),
        "redmarked": str(MANUSCRIPT_DIR / "eLife_V2_redmarked.docx"),
        "report": str(REPORT_DIR / "高阶拓扑论文V2升级完整报告.docx"),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
